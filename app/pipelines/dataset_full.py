"""Dataset stage sequencing keeps R, large frames, and Ollama separate."""

import gc
import json
from datetime import date
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from app.core.llm import unload_models
from app.models.analysis import AnalysisPlan
from app.pipelines.dataset import process_dataset_file
from app.pipelines.csv_repair import repair_csv_file
from app.pipelines.statistics import run_statistics, interpret_statistics


def export_dataset(parquet: Path, output_dir: Path) -> list[Path]:
    database = output_dir / "dataset.duckdb"
    xlsx = output_dir / "cleaned.xlsx"
    with duckdb.connect(str(database)) as db:
        db.execute("SET memory_limit='512MB'")
        db.execute("SET threads=2")
        db.execute(
            "CREATE TABLE cleaned AS SELECT * FROM read_parquet(?)", [str(parquet)]
        )
        cursor = db.execute("SELECT * FROM cleaned")
        headers = [c[0] for c in cursor.description]
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Cleaned data")

        def append_headers(target):
            cells = []
            for name in headers:
                cell = WriteOnlyCell(target, value=name)
                cell.data_type = "s"
                cells.append(cell)
            target.append(cells)

        append_headers(sheet)
        count = 1
        while rows := cursor.fetchmany(2000):
            for row in rows:
                if count >= 1_048_576:
                    sheet = workbook.create_sheet()
                    append_headers(sheet)
                    count = 1
                cells = []
                for value in row:
                    cell = WriteOnlyCell(sheet, value=value)
                    if isinstance(value, str):
                        cell.data_type = (
                            "s"  # Uploaded text can never become an Excel formula.
                        )
                    if isinstance(value, date):
                        cell.number_format = "dd/mm/yyyy"
                    cells.append(cell)
                sheet.append(cells)
                count += 1
        workbook.save(xlsx)
    return [database, xlsx]


def process_dataset(
    input_path: Path,
    output_dir: Path,
    *,
    date_order="dmy",
    trust_date_order=False,
    plan=None,
    interpret=True,
    record_mode="single_line",
    source_profile=None,
) -> list[Path]:
    unload_models()
    artifacts = clean_dataset(
        input_path,
        output_dir,
        date_order=date_order,
        trust_date_order=trust_date_order,
        record_mode=record_mode,
    )
    parquet = output_dir / "cleaned.parquet"
    artifacts += export_dataset(parquet, output_dir)
    gc.collect()
    counts = json.loads((output_dir / "repair_log.json").read_text())
    accepted = counts["valid_rows"] + counts["repaired_rows"]
    if accepted:
        statistics = run_statistics(
            parquet, output_dir, AnalysisPlan.model_validate(plan or {})
        )
        artifacts += [statistics] + [
            p
            for p in (output_dir / "analysis_request.json", output_dir / "r.log")
            if p.exists()
        ]
        gc.collect()
        if interpret:
            artifacts += interpret_statistics(statistics, output_dir)
            rejections = output_dir / "interpretation_rejections.jsonl"
            if rejections.exists():
                artifacts.append(rejections)
    summary = output_dir / "dataset_summary.json"
    summary.write_text(
        json.dumps(
            {
                "accepted_rows": accepted,
                "quarantined_rows": counts["ambiguous_rows"],
                "analysis_population": "Only accepted rows; review exclusion bias before interpreting results",
                "source_profile": source_profile,
                "interpretation_enabled": interpret,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    return [*artifacts, summary]


def clean_dataset(
    input_path: Path,
    output_dir: Path,
    *,
    date_order="dmy",
    trust_date_order=False,
    record_mode="single_line",
):
    """Repair and type data without loading R or Qwen."""
    repair = repair_csv_file(
        input_path,
        output_dir,
        date_order=date_order,
        trust_date_order=trust_date_order,
        record_mode=record_mode,
    )
    # The repair stage writes canonical DMY regardless of source convention.
    artifacts = process_dataset_file(repair.repaired_path, output_dir, date_order="dmy")
    artifacts += repair.artifacts
    report_path = output_dir / "dataset_report.json"
    report = json.loads(report_path.read_text())
    report.update(
        source_file=str(input_path.resolve()),
        input_date_order=date_order,
        trust_date_order=trust_date_order,
        headers=dict(zip(repair.original_headers, repair.normalized_headers)),
    )
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    return artifacts
